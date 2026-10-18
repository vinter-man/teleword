"""Exporting the user's words to xlsx / json / xml / csv."""

import csv
import io
import json
from dataclasses import asdict, dataclass
from enum import StrEnum
from xml.etree import ElementTree

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from sqlalchemy import func, select
from sqlalchemy.ext.asyncio import AsyncSession

from teleword.db.models import Example, Word


class Selection(StrEnum):
    ALL = "all"
    IMPORTANT = "important"  # rating above the user's average


class Order(StrEnum):
    ADDED = "added"
    ALPHABET = "alphabet"
    IMPORTANCE = "importance"  # highest rating first


class Format(StrEnum):
    XLSX = "xlsx"
    JSON = "json"
    XML = "xml"
    CSV = "csv"


@dataclass(frozen=True)
class Row:
    word_id: int
    word: str
    description: str
    example_id: int
    example: str


COLUMNS = ("word_id", "word", "description", "example_id", "example")


async def fetch_rows(session: AsyncSession, user_id: int, selection: Selection, order: Order) -> list[Row]:
    query = (
        select(Word.id, Word.word, Word.description, Example.id, Example.text)
        .join(Word.example)
        .where(Example.user_id == user_id)
    )
    if selection is Selection.IMPORTANT:
        average = (
            select(func.avg(Word.rating))
            .join(Word.example)
            .where(Example.user_id == user_id)
            .scalar_subquery()
        )
        query = query.where(Word.rating > average)

    query = query.order_by(
        {
            Order.ADDED: Word.id,
            Order.ALPHABET: func.lower(Word.word),
            Order.IMPORTANCE: Word.rating.desc(),
        }[order],
        Word.id,
    )
    return [Row(*r) for r in (await session.execute(query)).all()]


def filename(fmt: Format) -> str:
    return f"teleword_words.{fmt.value}"


def render(rows: list[Row], fmt: Format) -> bytes:
    return {
        Format.XLSX: _xlsx,
        Format.JSON: _json,
        Format.XML: _xml,
        Format.CSV: _csv,
    }[fmt](rows)


def _json(rows: list[Row]) -> bytes:
    return json.dumps([asdict(r) for r in rows], indent=4, ensure_ascii=False).encode()


def _xml(rows: list[Row]) -> bytes:
    root = ElementTree.Element("words")
    for row in rows:
        item = ElementTree.SubElement(root, "w")
        for column in COLUMNS:
            ElementTree.SubElement(item, column).text = str(getattr(row, column))
    return ElementTree.tostring(root, encoding="utf-8", xml_declaration=True)


def _csv(rows: list[Row]) -> bytes:
    buffer = io.StringIO()
    writer = csv.writer(buffer, quoting=csv.QUOTE_ALL)
    writer.writerow(COLUMNS)
    writer.writerows((r.word_id, r.word, r.description, r.example_id, r.example) for r in rows)
    # BOM so that Excel opens UTF-8 correctly
    return buffer.getvalue().encode("utf-8-sig")


_BG = PatternFill("solid", fgColor="1E252B")
_TEXT = Font(name="Segoe UI", size=14, color="B3BAC0")
_ACCENT = Font(name="Segoe UI", size=14, color="62ACBE")
_EDGE = Side(border_style="thin", color="0D0D0D")
_BORDER = Border(left=_EDGE, right=_EDGE, top=_EDGE, bottom=_EDGE)
_CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
_WIDTHS = {"A": 4, "B": 11, "C": 20, "D": 50, "E": 11, "F": 50}
_HEADERS = ("word id", "word", "description", "ex id", "example")


def _xlsx(rows: list[Row]) -> bytes:
    """A dark-themed table starting at B2, with ids and headers highlighted."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "words"
    for column, width in _WIDTHS.items():
        sheet.column_dimensions[column].width = width

    sheet.append([])
    sheet.append([None, *_HEADERS])
    for r in rows:
        sheet.append([None, r.word_id, r.word, r.description, r.example_id, r.example])
    last_row = sheet.max_row
    sheet.cell(row=last_row + 150, column=22, value="you found the easter egg :)")

    for row in sheet.iter_rows(min_row=1, max_row=sheet.max_row, max_col=22):
        for cell in row:
            cell.fill = _BG
            cell.font = _TEXT
            cell.alignment = _CENTER
            in_table = 2 <= cell.row <= last_row and 2 <= cell.column <= 6
            if in_table:
                cell.border = _BORDER
                if cell.row == 2 or cell.column in (2, 5):
                    cell.font = _ACCENT

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
